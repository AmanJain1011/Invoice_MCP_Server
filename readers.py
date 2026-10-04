import csv
from pathlib import Path

import pdfplumber
from openpyxl import load_workbook


DEVANAGARI_DIGITS = {chr(0x0966 + i): str(i) for i in range(10)}


def normalize_digits(text):
    for dev, asc in DEVANAGARI_DIGITS.items():
        text = text.replace(dev, asc)
    return text


def read_pdf(path):
    pages = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            pages.append(page.extract_text() or "")
    return "\n".join(pages)


def read_txt(path):
    return Path(path).read_text(encoding="utf-8")


def read_csv(path):
    with open(path, newline="", encoding="utf-8-sig") as f:
        return "\n".join(" | ".join(row) for row in csv.reader(f))


def read_excel(path):
    wb = load_workbook(path, data_only=True)
    lines = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            cells = ["" if c is None else str(c) for c in row]
            if any(cells):
                lines.append(" | ".join(cells))
    return "\n".join(lines)


READERS = {".pdf": read_pdf, ".txt": read_txt, ".csv": read_csv, ".xlsx": read_excel}


def read_document(file_path):
    ext = Path(file_path).suffix.lower()
    if ext not in READERS:
        raise ValueError(f"Unsupported file type: {ext}. Supported: {', '.join(READERS)}")
    return normalize_digits(READERS[ext](file_path)).strip()
